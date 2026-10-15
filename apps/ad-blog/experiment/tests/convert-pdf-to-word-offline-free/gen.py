"""PDF→Word 변환 테스트용 PDF 11개와 정답(truth.json) 생성.

본문 문장: Project Gutenberg 퍼블릭 도메인 책 3권
  https://www.gutenberg.org/cache/epub/1342/pg1342.txt  (Pride and Prejudice)
  https://www.gutenberg.org/cache/epub/84/pg84.txt      (Frankenstein)
  https://www.gutenberg.org/cache/epub/11/pg11.txt      (Alice's Adventures in Wonderland)
표·송장·이력서 항목은 이 스크립트가 만든 가상 값.

PDF 생성기
  - reportlab (문서 01~08)
  - python-docx로 만든 .docx → LibreOffice Writer PDF 내보내기 (문서 09~10, "Word에서 PDF로 저장"한 파일 역할)
  - 01을 200dpi 이미지로 만든 스캔 PDF (문서 11)

사용: python gen.py <gutenberg 폴더> <출력 폴더> <soffice 경로>
"""

import io
import json
import random
import re
import subprocess
import sys
from pathlib import Path

import fitz  # PyMuPDF
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches
from PIL import Image, ImageFilter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    Image as RLImage,
    ListFlowable,
    ListItem,
    PageTemplate,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

random.seed(20261015)
GUT, OUT, SOFFICE = Path(sys.argv[1]), Path(sys.argv[2]), sys.argv[3]
OUT.mkdir(parents=True, exist_ok=True)
SS = getSampleStyleSheet()
BODY = ParagraphStyle("body", parent=SS["Normal"], fontName="Helvetica", fontSize=11, leading=15, spaceAfter=8)


def paragraphs(book, n, skip=40):
    text = (GUT / book).read_text(encoding="utf-8")
    text = text.split("*** START")[1].split("*** END")[0]
    paras = [re.sub(r"\s+", " ", p).strip() for p in re.split(r"\n\s*\n", text)]
    paras = [p for p in paras if 45 <= len(p.split()) <= 140 and "_" not in p and "[" not in p]
    return paras[skip : skip + n]


POOL = paragraphs("pg1342.txt", 60) + paragraphs("pg84.txt", 60) + paragraphs("pg11.txt", 30, skip=10)
random.shuffle(POOL)


def take(n):
    return [POOL.pop() for _ in range(n)]


def esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


truth = {}


def record(name, title, paras, tables=(), images=0, extra=()):
    """정답: 문단(문장 단위 비교), 표(셀), 이미지 수, 기타 짧은 줄(제목·목록)."""
    truth[name] = {
        "title": title,
        "paragraphs": list(paras),
        "tables": [list(map(list, t)) for t in tables],
        "images": images,
        "lines": [title, *extra],
    }


def grid(data, header=True):
    t = Table(data, hAlign="LEFT")
    style = [("GRID", (0, 0), (-1, -1), 0.5, colors.grey), ("FONTSIZE", (0, 0), (-1, -1), 10)]
    if header:
        style += [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dde4ee")), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold")]
    t.setStyle(TableStyle(style))
    return t


def chart_png(seed, kind):
    r = random.Random(seed)
    fig, ax = plt.subplots(figsize=(5, 2.6), dpi=150)
    xs = ["Jan", "Feb", "Mar", "Apr", "May", "Jun"]
    ys = [r.randint(20, 90) for _ in xs]
    (ax.bar if kind == "bar" else ax.plot)(xs, ys, color="#3366aa")
    ax.set_ylabel("Orders")
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png")
    plt.close(fig)
    buf.seek(0)
    return buf


# 01 letter: 1쪽 편지
def doc_letter():
    paras = take(4)
    head = ["Northfield Garden Supply", "418 Orchard Lane, Springfield, IL 62704", "October 2, 2026", "Dear Ms. Bennet,"]
    story = [Paragraph(esc(h), BODY) for h in head] + [Paragraph(esc(p), BODY) for p in paras]
    story += [Paragraph("Sincerely,", BODY), Paragraph("Thomas Hale", BODY)]
    SimpleDocTemplate(str(OUT / "01-letter.pdf"), pagesize=letter).build(story)
    record("01-letter", head[0], paras, extra=head[1:] + ["Sincerely,", "Thomas Hale"])


# 02 report: 3쪽 보고서 (제목·소제목·글머리표·번호 목록)
def doc_report():
    paras = take(9)
    bullets = ["Review the supplier contracts before March", "Move archived files to the shared drive", "Schedule the annual safety inspection", "Update the customer contact list"]
    steps = ["Open the request form", "Fill in the project code", "Attach the signed quote", "Send it to the finance team"]
    story = [Paragraph("Quarterly Operations Report", SS["Title"]), Paragraph("1. Summary", SS["Heading2"])]
    story += [Paragraph(esc(p), BODY) for p in paras[:3]]
    story += [Paragraph("2. Action items", SS["Heading2"]), ListFlowable([ListItem(Paragraph(b, BODY)) for b in bullets], bulletType="bullet")]
    story += [Paragraph(esc(p), BODY) for p in paras[3:6]]
    story += [Paragraph("3. How to request a purchase", SS["Heading2"]), ListFlowable([ListItem(Paragraph(s, BODY)) for s in steps], bulletType="1")]
    story += [Paragraph(esc(p), BODY) for p in paras[6:]]
    SimpleDocTemplate(str(OUT / "02-report.pdf"), pagesize=letter).build(story)
    record("02-report", "Quarterly Operations Report", paras, extra=["1. Summary", "2. Action items", "3. How to request a purchase", *bullets, *steps])


# 03 invoice: 표 1개 (9행 4열)
def doc_invoice():
    items = ["Garden hose 50 ft", "Pruning shears", "Potting soil 40 lb", "Cedar planter box", "Drip irrigation kit", "Leaf rake", "Seed starter tray"]
    rows = [["Item", "Qty", "Unit price", "Amount"]]
    total = 0.0
    for it in items:
        q, p = random.randint(1, 9), round(random.uniform(4, 120), 2)
        total += q * p
        rows.append([it, str(q), f"${p:,.2f}", f"${q * p:,.2f}"])
    rows.append(["Total", "", "", f"${total:,.2f}"])
    paras = take(1)
    head = ["Invoice #10482", "Billed to: Riverside Community Center", "Due date: November 1, 2026"]
    story = [Paragraph(head[0], SS["Title"])] + [Paragraph(h, BODY) for h in head[1:]] + [Spacer(1, 10), grid(rows), Spacer(1, 12), Paragraph(esc(paras[0]), BODY)]
    SimpleDocTemplate(str(OUT / "03-invoice.pdf"), pagesize=letter).build(story)
    record("03-invoice", head[0], paras, tables=[rows], extra=head[1:])


# 04 schedule: 표 2개 + 본문
def doc_schedule():
    t1 = [["Day", "Time", "Session", "Room"]] + [[d, f"{h}:00", s, r] for d, h, s, r in [("Monday", 9, "Welcome and setup", "A1"), ("Monday", 13, "Spreadsheet basics", "B2"), ("Tuesday", 10, "Email safety", "A1"), ("Tuesday", 14, "Video calls", "C3"), ("Wednesday", 11, "Backing up photos", "B2")]]
    t2 = [["Name", "Role", "Extension"], ["Maria Lopez", "Coordinator", "2201"], ["James Carter", "Instructor", "2214"], ["Aisha Khan", "Volunteer", "2230"]]
    paras = take(3)
    story = [Paragraph("Community Tech Week Schedule", SS["Title"]), Paragraph(esc(paras[0]), BODY), grid(t1), Spacer(1, 14), Paragraph(esc(paras[1]), BODY), grid(t2), Spacer(1, 14), Paragraph(esc(paras[2]), BODY)]
    SimpleDocTemplate(str(OUT / "04-schedule.pdf"), pagesize=letter).build(story)
    record("04-schedule", "Community Tech Week Schedule", paras, tables=[t1, t2])


# 05 newsletter: 2단 편집, 2쪽
def doc_newsletter():
    paras = take(10)
    doc = BaseDocTemplate(str(OUT / "05-newsletter.pdf"), pagesize=letter)
    w, h = letter
    m, g = 0.75 * inch, 0.3 * inch
    cw = (w - 2 * m - g) / 2
    frames = [Frame(m, m, cw, h - 2 * m, id="c1"), Frame(m + cw + g, m, cw, h - 2 * m, id="c2")]
    doc.addPageTemplates([PageTemplate(id="two", frames=frames)])
    story = [Paragraph("The Maple Street Monthly", SS["Heading1"]), Paragraph("Neighborhood news", SS["Heading3"])]
    story += [Paragraph(esc(p), BODY) for p in paras]
    doc.build(story)
    record("05-newsletter", "The Maple Street Monthly", paras, extra=["Neighborhood news"])


# 06 images: 본문 + 차트 이미지 2개 + 캡션
def doc_images():
    paras = take(4)
    caps = ["Figure 1. Monthly orders, first half of the year", "Figure 2. Orders trend by month"]
    story = [Paragraph("Spring Sales Review", SS["Title"]), Paragraph(esc(paras[0]), BODY), RLImage(chart_png(1, "bar"), width=5 * inch, height=2.6 * inch), Paragraph(caps[0], BODY), Paragraph(esc(paras[1]), BODY), Paragraph(esc(paras[2]), BODY), RLImage(chart_png(2, "line"), width=5 * inch, height=2.6 * inch), Paragraph(caps[1], BODY), Paragraph(esc(paras[3]), BODY)]
    SimpleDocTemplate(str(OUT / "06-images.pdf"), pagesize=letter).build(story)
    record("06-images", "Spring Sales Review", paras, images=2, extra=caps)


# 07 resume: 이력서 (테두리 없는 2열 표 + 글머리표)
def doc_resume():
    bullets = ["Managed a team of six customer support agents", "Cut average reply time from two days to five hours", "Trained new staff on the help desk software", "Wrote the internal guide for refund requests"]
    side = [["Email", "jordan.reyes@example.com"], ["Phone", "(555) 014-2290"], ["City", "Portland, OR"]]
    paras = take(1)
    story = [Paragraph("Jordan Reyes", SS["Title"]), Paragraph("Customer Support Lead", SS["Heading3"])]
    t = Table(side, hAlign="LEFT", colWidths=[1.2 * inch, 3 * inch])
    t.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 10)]))
    story += [t, Paragraph("Summary", SS["Heading2"]), Paragraph(esc(paras[0]), BODY), Paragraph("Experience", SS["Heading2"]), Paragraph("Support Lead, Brightline Outfitters, 2021 to 2026", BODY), ListFlowable([ListItem(Paragraph(b, BODY)) for b in bullets], bulletType="bullet")]
    SimpleDocTemplate(str(OUT / "07-resume.pdf"), pagesize=letter).build(story)
    record("07-resume", "Jordan Reyes", paras, tables=[side], extra=["Customer Support Lead", "Summary", "Experience", "Support Lead, Brightline Outfitters, 2021 to 2026", *bullets])


# 08 머리글·바닥글·쪽 번호가 있는 3쪽 문서
def doc_headerfooter():
    paras = take(10)

    def deco(c, d):
        c.saveState()
        c.setFont("Helvetica", 9)
        c.drawString(inch, letter[1] - 0.6 * inch, "Riverside Library Policy Manual")
        c.drawRightString(letter[0] - inch, 0.6 * inch, f"Page {d.page}")
        c.restoreState()

    story = [Paragraph("Borrowing and Returns Policy", SS["Title"])] + [Paragraph(esc(p), BODY) for p in paras]
    SimpleDocTemplate(str(OUT / "08-header-footer.pdf"), pagesize=letter).build(story, onFirstPage=deco, onLaterPages=deco)
    record("08-header-footer", "Borrowing and Returns Policy", paras)


def to_pdf_with_writer(docx_path):
    subprocess.run([SOFFICE, "--headless", "--convert-to", "pdf", "--outdir", str(OUT), str(docx_path)], check=True, capture_output=True)


# 09 Word 문서(제목 스타일·표·글머리표)를 PDF로 내보낸 파일
def doc_word_report():
    paras = take(5)
    rows = [["Region", "Q1", "Q2", "Q3"], ["North", "1,240", "1,310", "1,402"], ["South", "980", "1,105", "1,190"], ["West", "1,515", "1,488", "1,620"]]
    bullets = ["Hire two more drivers", "Open the new warehouse in May", "Review delivery routes monthly"]
    d = Document()
    d.add_heading("Regional Delivery Plan", 0)
    d.add_heading("Background", 1)
    for p in paras[:2]:
        d.add_paragraph(p)
    d.add_heading("Deliveries by region", 1)
    t = d.add_table(rows=len(rows), cols=4)
    t.style = "Table Grid"
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            t.cell(i, j).text = v
    d.add_paragraph(paras[2])
    d.add_heading("Next steps", 1)
    for b in bullets:
        d.add_paragraph(b, style="List Bullet")
    for p in paras[3:]:
        d.add_paragraph(p)
    src = OUT / "src" / "09-word-report.docx"
    src.parent.mkdir(exist_ok=True)
    d.save(src)
    to_pdf_with_writer(src)
    record("09-word-report", "Regional Delivery Plan", paras, tables=[rows], extra=["Background", "Deliveries by region", "Next steps", *bullets])


# 10 Word 문서(이미지·번호 목록)를 PDF로 내보낸 파일
def doc_word_image():
    paras = take(4)
    steps = ["Unplug the router for thirty seconds", "Plug it back in and wait for the lights", "Restart your computer", "Run the speed test again"]
    d = Document()
    d.add_heading("Home Network Checkup", 0)
    d.add_paragraph(paras[0])
    d.add_picture(chart_png(3, "bar"), width=Inches(5))
    d.add_paragraph("Figure 1. Download speed by day")
    d.add_paragraph(paras[1])
    for s in steps:
        d.add_paragraph(s, style="List Number")
    for p in paras[2:]:
        d.add_paragraph(p)
    src = OUT / "src" / "10-word-image.docx"
    src.parent.mkdir(exist_ok=True)
    d.save(src)
    to_pdf_with_writer(src)
    record("10-word-image", "Home Network Checkup", paras, images=1, extra=["Figure 1. Download speed by day", *steps])


# 11 스캔 PDF: 01을 200dpi 이미지로 (약간 기울임·흐림)
def doc_scanned():
    src = fitz.open(OUT / "01-letter.pdf")
    out = fitz.open()
    for page in src:
        pix = page.get_pixmap(dpi=200, colorspace=fitz.csGRAY)
        img = Image.frombytes("L", (pix.width, pix.height), pix.samples).rotate(0.6, fillcolor=255, resample=Image.BICUBIC).filter(ImageFilter.GaussianBlur(0.6))
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=80)
        p = out.new_page(width=page.rect.width, height=page.rect.height)
        p.insert_image(p.rect, stream=buf.getvalue())
    out.save(OUT / "11-scanned.pdf")
    t = dict(truth["01-letter"])
    truth["11-scanned"] = t


for f in [doc_letter, doc_report, doc_invoice, doc_schedule, doc_newsletter, doc_images, doc_resume, doc_headerfooter, doc_word_report, doc_word_image, doc_scanned]:
    f()

(OUT / "truth.json").write_text(json.dumps(truth, indent=1, ensure_ascii=False))
for pdf in sorted(OUT.glob("*.pdf")):
    print(pdf.name, len(fitz.open(pdf)), "pages")
