"""예시 Gem과 지식 파일 생성 (테스트 입력, 직접 작성한 내용)."""
import json
from pathlib import Path
from docx import Document

root = Path('sample')
k = root / 'knowledge'
k.mkdir(parents=True, exist_ok=True)
gem = {
    'name': 'Weekly Newsletter Writer',
    'description': 'Drafts our weekly community newsletter in a friendly, plain-English voice. Use when writing or editing the newsletter, summarizing club news, or turning event notes into a short update.',
    'instructions': (
        'You write the weekly newsletter for a neighborhood gardening club.\n\n'
        '1. Start with a two-sentence summary of the week.\n'
        '2. Use the sections: Upcoming events, Garden tips, Volunteer needs.\n'
        '3. Keep it under 300 words and use plain English.\n'
        '4. Follow the tone rules in the style guide file.\n\n'
        'Common mistakes to avoid: do not invent event dates; if a date is missing, ask for it.'
    ),
}
(root / 'gem.json').write_text(json.dumps(gem, indent=2))
d = Document()
d.add_heading('Newsletter style guide', 1)
d.add_paragraph('Use short sentences. Address readers as "you". Never use exclamation marks in headings.')
d.save(k / 'style-guide.docx')
(k / 'event-calendar.csv').write_text('date,event\n2026-10-24,Bulb planting day\n2026-11-07,Seed swap\n')
(k / 'past-issue.md').write_text('# Issue 41\n\nThis week we harvested the last tomatoes...\n')
(k / '.DS_Store').write_bytes(b'\x00\x00\x00\x01Bud1' + b'\x00' * 100)
print('sample ready')
